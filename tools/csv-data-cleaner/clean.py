import argparse
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None


ENCODINGS = ("utf-8-sig", "utf-8", "cp932", "gb18030")
DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
    "%Y年%m月%d日",
    "%m/%d/%Y",
    "%d/%m/%Y",
)


def detect_encoding(data):
    for encoding in ENCODINGS:
        try:
            data.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    raise UnicodeError("encoding could not be detected")


def detect_delimiter(text):
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def text_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def normalize_date(value):
    value = value.strip()
    if not value:
        return value
    for date_format in DATE_FORMATS:
        try:
            return datetime.strptime(value, date_format).date().isoformat()
        except ValueError:
            continue
    return value


def normalize_amount(value):
    original = value.strip()
    if not original:
        return original
    negative = original.startswith("(") and original.endswith(")")
    cleaned = re.sub(r"(?:JPY|USD|EUR|CNY|HKD)", "", original, flags=re.IGNORECASE)
    cleaned = re.sub(r"[¥￥$€£,\s]", "", cleaned)
    if negative:
        cleaned = "-" + cleaned[1:-1]
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        return original
    if number == number.to_integral():
        return str(number.quantize(Decimal("1")))
    return format(number.normalize(), "f")


def clean_table(header, rows, date_columns=None, amount_columns=None, keep_empty=False, keep_duplicates=False):
    date_columns = set(date_columns or [])
    amount_columns = set(amount_columns or [])
    width = max([len(header), *(len(row) for row in rows)] or [0])
    header = [text_value(value).strip() for value in header]
    header.extend(f"column_{index + 1}" for index in range(len(header), width))

    cleaned = []
    seen = set()
    removed_empty = 0
    removed_duplicates = 0
    for source in rows:
        row = [text_value(value).strip() for value in source]
        row.extend("" for _ in range(width - len(row)))
        for index, name in enumerate(header):
            if name in date_columns:
                row[index] = normalize_date(row[index])
            if name in amount_columns:
                row[index] = normalize_amount(row[index])
        key = tuple(row)
        if not keep_empty and not any(row):
            removed_empty += 1
            continue
        if not keep_duplicates and key in seen:
            removed_duplicates += 1
            continue
        seen.add(key)
        cleaned.append(row)
    return header, cleaned, removed_empty, removed_duplicates


def read_csv(filename, encoding="auto", delimiter="auto"):
    data = Path(filename).read_bytes()
    selected_encoding = detect_encoding(data) if encoding == "auto" else encoding
    text = data.decode(selected_encoding)
    selected_delimiter = detect_delimiter(text) if delimiter == "auto" else delimiter
    rows = list(csv.reader(io.StringIO(text), delimiter=selected_delimiter))
    if not rows:
        return [], [], selected_delimiter, selected_encoding
    return rows[0], rows[1:], selected_delimiter, selected_encoding


def read_xlsx(filename):
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for xlsx files")
    workbook = load_workbook(filename, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        return [], []
    return rows[0], rows[1:]


def write_csv(filename, header, rows, delimiter=","):
    with Path(filename).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter=delimiter)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(filename, header, rows):
    if Workbook is None:
        raise RuntimeError("openpyxl is required for xlsx files")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(filename)


def output_path(filename, output_dir):
    source = Path(filename)
    suffix = source.suffix.lower()
    return Path(output_dir) / f"{source.stem}_clean{suffix}"


def clean_file(filename, output_dir, encoding="auto", delimiter="auto", date_columns=None, amount_columns=None, keep_empty=False, keep_duplicates=False):
    source = Path(filename)
    if source.suffix.lower() == ".xlsx":
        header, rows = read_xlsx(source)
        used_delimiter = ","
    else:
        header, rows, used_delimiter, _ = read_csv(source, encoding, delimiter)

    header, rows, empty_count, duplicate_count = clean_table(
        header,
        rows,
        date_columns=date_columns,
        amount_columns=amount_columns,
        keep_empty=keep_empty,
        keep_duplicates=keep_duplicates,
    )
    destination = output_path(source, output_dir)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.suffix.lower() == ".xlsx":
        write_xlsx(destination, header, rows)
    else:
        write_csv(destination, header, rows, used_delimiter)
    return {
        "input": str(source),
        "output": str(destination),
        "rows": len(rows),
        "empty_removed": empty_count,
        "duplicates_removed": duplicate_count,
    }


def column_list(value):
    return [item.strip() for item in value.split(",") if item.strip()]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+")
    parser.add_argument("--output-dir", default="cleaned")
    parser.add_argument("--encoding", default="auto")
    parser.add_argument("--delimiter", default="auto")
    parser.add_argument("--date-columns", type=column_list, default=[])
    parser.add_argument("--amount-columns", type=column_list, default=[])
    parser.add_argument("--keep-empty", action="store_true")
    parser.add_argument("--keep-duplicates", action="store_true")
    args = parser.parse_args()

    failed = False
    for filename in args.files:
        try:
            result = clean_file(
                filename,
                args.output_dir,
                encoding=args.encoding,
                delimiter=args.delimiter,
                date_columns=args.date_columns,
                amount_columns=args.amount_columns,
                keep_empty=args.keep_empty,
                keep_duplicates=args.keep_duplicates,
            )
            print(
                f"{result['output']}  rows={result['rows']} "
                f"empty={result['empty_removed']} duplicates={result['duplicates_removed']}"
            )
        except Exception as error:
            failed = True
            print(f"{filename}: {error}")
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
